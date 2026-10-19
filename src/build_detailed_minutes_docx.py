#!/usr/bin/env python3
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("音频转录工具/家长因孩子被打追责对方_详细谈话纪要.docx")

BLACK = RGBColor(0, 0, 0)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(88, 88, 88)
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "EAF2F8"


def set_run_font(run, size=11, bold=None, color=BLACK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def para_spacing(p, before=0, after=6, line=1.1):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=90, bottom=90, start=120, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        el = margins.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            margins.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cell.paragraphs:
                para_spacing(p, after=3, line=1.08)


def add_p(doc, text, size=11, bold=False, color=BLACK, before=0, after=6):
    p = doc.add_paragraph()
    para_spacing(p, before=before, after=after)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    para_spacing(p, before=16 if level == 1 else 10, after=8 if level == 1 else 5)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else DARK_BLUE)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    para_spacing(p, after=5, line=1.12)
    r = p.add_run(text)
    set_run_font(r, size=10.8)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    para_spacing(p, after=5, line=1.12)
    r = p.add_run(text)
    set_run_font(r, size=10.8)
    return p


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    for style_name in ["Normal", "List Bullet", "List Bullet 2", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.line_spacing = 1.12
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    para_spacing(header, after=0)
    set_run_font(header.add_run("详细谈话纪要 | 家长因孩子被打追责对方"), size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para_spacing(footer, after=0)
    set_run_font(footer.add_run("根据自动转写整理，关键事实请结合录音、监控及正式笔录复核"), size=8.5, color=MUTED)


def metadata_table(doc):
    data = [
        ("整理对象", "幼儿园冲突处理会议录音"),
        ("录音文件", "/Users/linchenhao/Downloads/家长因孩子被打追责对方.m4a"),
        ("录音时长", "约 87 分钟"),
        ("文档性质", "详细谈话内容整理版，非公安笔录、非逐字稿"),
        ("整理口径", "按时间顺序梳理各方主要发言、追问、解释、程序说明和明确表态。录音中有多人交叉说话，个别细节需回听核实。"),
    ]
    table = doc.add_table(rows=len(data), cols=2)
    table.style = "Table Grid"
    set_cell_margins(table)
    set_table_width(table, [1500, 8300])
    for row, (k, v) in zip(table.rows, data):
        shade(row.cells[0], LIGHT_GRAY)
        for text, cell, bold in [(k, row.cells[0], True), (v, row.cells[1], False)]:
            p = cell.paragraphs[0]
            p.text = ""
            r = p.add_run(text)
            set_run_font(r, size=10.5, bold=bold)


def add_timeline_table(doc, blocks):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_cell_margins(table, top=100, bottom=100, start=120, end=120)
    set_table_width(table, [1200, 1600, 7000])
    headers = ["时间段", "主要发言方", "详细谈话内容整理"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        set_run_font(p.add_run(h), size=10.5, bold=True)

    for time, speaker, points in blocks:
        cells = table.add_row().cells
        for idx, value in enumerate([time, speaker]):
            p = cells[idx].paragraphs[0]
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(value), size=10.2, bold=(idx == 0), color=DARK_BLUE if idx == 0 else BLACK)
        p = cells[2].paragraphs[0]
        p.text = ""
        for n, point in enumerate(points, start=1):
            if n > 1:
                p = cells[2].add_paragraph()
            para_spacing(p, after=3, line=1.08)
            r = p.add_run(f"{n}. {point}")
            set_run_font(r, size=10.2)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    add_p(doc, "详细谈话纪要", size=23, bold=True, before=10, after=2)
    add_p(doc, "家长因孩子被打追责对方会议录音整理版", size=13.5, color=MUTED, after=12)
    metadata_table(doc)

    add_heading(doc, "一、阅读说明", 1)
    for item in [
        "本文件比简版会议纪要更侧重“谈话内容”，按录音时间顺序记录各方发言推进。",
        "由于自动转写中存在错字、漏字和多人同时说话，本文件采用整理表述；涉及时间、姓名、动作细节和法律结论的内容，建议以原录音、监控、医院材料、公安笔录为准。",
        "文中使用“被打方家长”“涉事家长”“园方”“公安/社区”“家委/律师代表”等角色称呼，避免将自动识别错误的人名写成事实。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "二、核心脉络", 1)
    for item in [
        "会议起初由被打方家长情绪化质问涉事家长，重点围绕“为什么不找老师/家长、为什么直接找幼儿园孩子、为什么动手打孩子、为什么事后离开”。",
        "涉事家长承认行为不当并道歉，但解释原本是想了解孩子之间此前矛盾、让两个孩子说清楚；被打方认为该解释不能成立。",
        "公安/社区和园方陆续介入，明确涉事家长直接找孩子并动手是不当行为，同时要求各方回到孩子伤情、心理状态和依法处理流程上。",
        "园方承认管理存在漏洞，说明事发后先安抚孩子、后通知家长，愿意配合报警、监控、上报、心理支持和整改。",
        "被打方家长最终明确：可以接受对方道歉这一事实，但不接受调解或谅解，要求依法、秉公处理，不主张额外加码或私下索赔。",
    ]:
        add_number(doc, item)

    add_heading(doc, "三、按时间顺序的详细谈话内容", 1)
    blocks = [
        ("00:00-00:06", "被打方家长", [
            "开场即质问涉事家长：既然前一晚已经与配偶商量要到幼儿园了解情况，为什么不找老师、不找家长，而是直接去找幼儿园孩子。",
            "反复表达自己看到孩子被打后情绪失控、整夜未眠，但强调自己并未去找对方孩子发泄。",
            "指出事发后对方没有留下来说明情况，老师和对方联系也不顺畅，导致自己最终报警。",
            "强调自己来会议不是单纯接受道歉，而是要当面追问涉事家长为什么可以对幼儿园孩子动手。",
        ]),
        ("00:06-00:09", "公安/社区、被打方家长", [
            "主持方表示理解父母看到孩子受伤后的情绪，希望各方把事情处理清楚；若协商不成则依法走程序。",
            "被打方回应称自己没有去找对方孩子，核心问题是两个成年人直接面对一个幼儿，本身就是不对等和不安全的。",
            "被打方质疑园方现场敏感性不足：两个不常来的家长突然找孩子，老师应当警觉并及时介入。",
            "被打方描述孩子被打后没有立即哭，认为孩子可能是被吓懵，而不是事情不严重。",
        ]),
        ("00:09-00:12", "被打方家长", [
            "根据老师描述复盘现场：老师刚要询问情况，涉事家长突然上前打了一巴掌，老师随后把孩子拉开护在身后。",
            "被打方认为涉事家长当时已有充分情绪积累，并非完全没有准备的一瞬间误动作。",
            "被打方提到孩子后来表示“事情过去了/原谅了”，但家长判断孩子更多是害怕爸爸骂人或被警察带走，并非真正不受影响。",
            "被打方再次强调自己愤怒但不会伤害对方孩子，认为成年人应有基本边界。",
        ]),
        ("00:12-00:15", "园方老师、涉事家长、被打方家长", [
            "园方老师表示，涉事家长动作太快，自己没有来得及完全拦住，随后立即把孩子拉开。",
            "涉事家长解释本意是想了解孩子之间的情况、让孩子把事情说清楚。",
            "被打方反问：如果只是了解情况，为什么不让老师主持、不等孩子讲完，为什么在孩子说话之前就动手。",
            "涉事家长道歉；被打方明确回应，道歉是应该的，但自己不接受把事情到此为止。",
        ]),
        ("00:15-00:18", "被打方家长、涉事家长", [
            "谈话转入此前孩子之间可能存在的矛盾，录音中提到过去发生过若干孩子间冲突，例如跳到身上、固体胶弄到嘴里等说法，具体需进一步核实。",
            "被打方表示，自己也曾因孩子被伤害而生气，但仍会先通过老师、家长群或园方渠道了解，而不是直接找孩子对质。",
            "涉事家长说明前一晚从自己孩子处听到情况，因此决定第二天到学校了解并沟通。",
            "被打方认为涉事家长即便有委屈，也有充分时间冷静和选择正确渠道，不能以“了解情况”为由对孩子施压或动手。",
        ]),
        ("00:18-00:24", "被打方家长、涉事家长", [
            "涉事家长再次表达歉意。被打方表示可以听到道歉，但不代表接受谅解或调解。",
            "被打方举例称，即便自己听说孩子被其他孩子多次欺负，也会先判断事情力度、孩子陈述是否完整、自己孩子是否也有责任。",
            "被打方强调，如果实在咽不下这口气，可以找老师、园方或对方家长，绝不能直接找幼儿园孩子。",
            "被打方认为孩子之间的矛盾应由老师和成年人之间的正式沟通解决，不能把压力转移给幼儿。",
        ]),
        ("00:24-00:30", "公安/社区、园方、被打方家长", [
            "主持方明确指出，涉事家长直接找孩子并动手是错误的；孩子在幼儿园打闹应首先相信老师和园方处理。",
            "主持方批评涉事家长作为父母给孩子做了不良示范，也提醒被打方家长尽量避免在自己孩子面前持续强化负面情绪。",
            "园方/主持方解释，园方未第一时间通知被打方家长，部分考虑是担心双方家长冲突进一步升级。",
            "被打方家长对“让涉事家长先离开”和“通知不及时”仍表示不能理解，认为这是园方管理和处置问题。",
            "讨论开始转向孩子身体伤情和心理状态，主持方提醒后续应优先关注孩子情况。",
        ]),
        ("00:30-00:36", "公安、被打方家长", [
            "被打方不接受“目前没大问题”等轻描淡写表达，认为这会淡化孩子受到的伤害。",
            "公安说明后续需提供医院就诊记录、检查材料，并可能向孩子、老师、医生等了解情况。",
            "公安解释伤情鉴定程序：是否需要做公安司法伤情鉴定由家长考虑；伤情等级会影响后续处理路径。",
            "公安说明若达到轻伤标准可能进入刑事程序；若未达到轻伤，通常按治安管理程序依法处理。",
            "公安表示可以组织调解，但调解必须尊重双方意愿，不是强制要求受害方接受。",
        ]),
        ("00:36-00:39", "被打方家长", [
            "被打方表示，如果孩子几天后没有明显身体异常，自己未必坚持走复杂伤情鉴定流程。",
            "同时强调，如果孩子后续出现症状或医生建议进一步检查，就会依法走相应程序。",
            "被打方说明自己并非为了额外“整”对方，而是要求成年人为已经发生的动手行为承担后果。",
            "被打方认为处罚本身就是对其行为的基本回应，不是自己额外施压。",
        ]),
        ("00:39-00:45", "被打方家长", [
            "被打方用多个假设场景说明，涉事家长容易因孩子受委屈而冲动，未来如果孩子再遇到更严重冲突，可能产生更大风险。",
            "被打方表示自己一整晚反复思考是否接受调解，最终认为不接受调解不是一时冲动。",
            "被打方提出，受害者没有义务因为对方道歉就原谅；如果接受“道歉就算了”，反而可能变成对受害方的道德绑架。",
            "被打方提出依法处罚的诉求，表述为“该怎么处理就怎么处理”，不要求额外加码。",
        ]),
        ("00:45-00:48", "被打方家长", [
            "被打方继续解释，不接受调解不等于“得理不饶人”，而是作为受害方要求公平处理。",
            "被打方提到孩子临近毕业，原本期待幼儿园活动，现在家长对继续送孩子上学产生顾虑。",
            "被打方强调自己没有提出额外赔偿或私下索取，只希望通过依法处理给孩子和自己一个交代。",
        ]),
        ("00:48-00:54", "园方负责人", [
            "园方说明，事发大约在 8:52 左右；老师发现后第一时间处理孩子安全和现场秩序。",
            "园方表示，当时不止两个当事孩子，现场还有其他孩子看到，因此需要先安抚孩子们的情绪。",
            "园方称涉事家长当场已有道歉或交流，随后园方让涉事家长先回去，但要求其随时配合后续处理。",
            "园方解释，学校属于重点管理场所，若家长冲突影响正常教学秩序，事情可能进一步升级，因此先避免冲突扩大。",
            "园方说明，被打孩子起初不愿联系父母，担心被骂；老师持续安抚和做工作，约 10:30 左右才联系家长。",
            "园方承认管理存在漏洞，不推卸、不逃避；表示支持报警、调取监控、依法追责，并已向教育局、社区及全体家长说明情况。",
        ]),
        ("00:54-00:57", "家委/律师代表、园方、涉事家长", [
            "家委/律师代表表示，公安已经介入，后续应在查明事实基础上依法、公正处理。",
            "代表认为调解或秉公处理都可以是解决方式，但必须尊重各方自愿，尤其是被打方是否接受调解。",
            "园方提出可为孩子安排心理支持，并倾向优先使用孩子熟悉的心理相关人员，避免陌生干预带来新的压力。",
            "涉事家长再次表达歉意，同时提出希望不要因为自己的错误影响自己孩子未来的学习和成长。",
        ]),
        ("00:57-01:03", "被打方家长、涉事家长", [
            "被打方正面回应涉事家长“给余地”的诉求，认为不能把涉事家长自己的后果变成受害方的道德负担。",
            "被打方表示理解父母保护孩子的冲动，但这种冲动不能成为对别人孩子动手的理由。",
            "被打方认为涉事家长经过一晚情绪积累后到学校找孩子，说明并非完全没有时间冷静。",
            "被打方也谈到自己知道自己脾气急，因此来处理事情时会安排家人或老师在旁边，避免自己失控；但自己从未想过伤害对方孩子。",
        ]),
        ("01:03-01:09", "涉事家长、被打方家长", [
            "涉事家长解释自己事后在公司处理客户事务，后来看到老师或公安电话后赶来。",
            "被打方追问，如果真心意识到错误，为什么没有更主动地通过老师、园方或家长表达歉意和说明情况。",
            "被打方表示，如果涉事家长事发后留下来、第一时间当面道歉，自己可能不会走到报警这一步；但这只是其情绪判断，并不改变依法处理的立场。",
            "被打方再次强调自己虽然愤怒，却没有找对方孩子，也没有把情绪转移到幼儿身上。",
        ]),
        ("01:09-01:15", "被打方家长", [
            "被打方认为，涉事家长后来低姿态道歉、表现后悔，不等于受害方必须出具谅解或放弃追责。",
            "被打方反复说明自己不是要陷害对方，也不索要额外赔偿，只是要求对方为自己已经做出的行为承担应有后果。",
            "被打方担心外界可能在事后同情施暴者、反过来指责受害方“得理不饶人”，因此反复解释自己不接受调解的逻辑。",
            "被打方表示，这件事对孩子和家庭造成心理影响，后续是否恢复需要继续观察。",
        ]),
        ("01:15-01:18", "公安/社区、被打方家长", [
            "主持方总结：被打方情绪可以理解，打人方道歉态度也可以看到，但是否接受道歉、调解或谅解完全由被打方自愿决定。",
            "主持方建议先把孩子放在第一位，观察三天左右，看身体和心理状态再决定后续。",
            "主持方表示园方也已经介绍了紧急措施和后续安排，希望家长理解并继续沟通。",
        ]),
        ("01:18-01:24", "公安/社区、被打方家长、涉事家长", [
            "主持方再次说明公安流程和可能面临的处罚，同时提醒双方做好孩子心理引导。",
            "涉事家长提出是否可以去看孩子或再次道歉；被打方拒绝，理由是希望淡化事件，不在孩子面前再次强化记忆。",
            "被打方表示，孩子内心并没有真正过去，更多是害怕爸爸骂人或被带走；因此不能简单理解为孩子说原谅就代表无影响。",
            "被打方明确表示：道歉可以接受，但事情仍然要秉公办理，这是给孩子和自己一个交代。",
            "被打方继续提醒涉事家长，不要把依法处理产生的后果理解为受害方造成，根源是涉事家长自己的行为。",
        ]),
        ("01:24-01:27", "各方收尾", [
            "被打方再次明确：可以接受对方道歉，但不接受将事情到此为止；后续该怎么处理就怎么处理。",
            "各方围绕后续联系、公安处理时间、是否再沟通等事项做简短确认。",
            "会议逐步收尾。自动转写最后数秒出现明显误识别内容，已忽略。",
        ]),
    ]
    add_timeline_table(doc, blocks)

    add_heading(doc, "四、各方明确表态整理", 1)
    statements = [
        ("被打方家长", [
            "接受“道歉是应当的”，但不接受调解或谅解。",
            "要求依法、秉公处理，不要求额外加码，也不以私下赔偿作为核心诉求。",
            "不希望涉事家长再见孩子，以免强化孩子记忆或造成二次刺激。",
            "会根据孩子后续身体和心理情况决定是否进一步就医或申请伤情鉴定。",
        ]),
        ("涉事家长", [
            "承认事情做错，向被打孩子和家长道歉。",
            "解释最初想法是了解孩子之间的矛盾，但承认自己情绪失控并动手。",
            "希望被打方考虑自己孩子的后续影响，留出一定余地。",
        ]),
        ("园方", [
            "承认事件发生在园内，管理存在漏洞。",
            "支持报警、查看监控、依法追责，并配合公安和主管部门。",
            "已向教育局、社区及全体家长说明情况，愿意向被打方郑重道歉。",
            "愿意根据家长意愿提供心理支持，并复盘完善制度。",
        ]),
        ("公安/社区", [
            "明确涉事家长直接找幼儿园孩子并动手是不当行为。",
            "后续根据医院材料、询问情况、伤情判断依法处理。",
            "调解可以组织，但是否接受完全尊重双方自愿。",
        ]),
    ]
    for role, items in statements:
        add_heading(doc, role, 2)
        for item in items:
            add_bullet(doc, item)

    add_heading(doc, "五、后续事项", 1)
    for item in [
        "被打方家长继续观察孩子身体、心理状态，保存医院记录和相关凭证。",
        "被打方家长根据孩子情况决定是否申请公安司法伤情鉴定。",
        "园方保存并按程序提供监控、现场情况说明、老师沟通记录等材料。",
        "公安继续收集证据、询问相关人员，并依法作出处置。",
        "园方复盘家长入园、家长接触幼儿、突发冲突上报、家校通知等制度漏洞。",
        "未经被打方同意，不安排涉事家长再次接触被打孩子。",
    ]:
        add_number(doc, item)

    add_heading(doc, "六、备注", 1)
    add_p(doc, "本文件是根据自动转写和录音内容进行的详细整理，不替代公安笔录、医疗诊断、伤情鉴定、监控资料或园方正式调查报告。涉及责任认定、处罚结果、伤情等级等事项，应以公安机关、医疗机构和相关主管部门正式文件为准。")

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
