#!/usr/bin/env python3
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("meeting_minutes_audio/家长因孩子被打追责对方_会议纪要.docx")

BLACK = RGBColor(0, 0, 0)
HEADING_BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_GRAY = "F2F4F7"
BORDER = "D9E2F3"


def set_run_font(run, size=None, bold=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.10):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(table, top=80, bottom=80, start=120, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for side, value in [("top", top), ("bottom", bottom), ("start", start), ("end", end)]:
        el = margins.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            margins.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                set_paragraph_spacing(p, after=2, line=1.10)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=16 if level == 1 else 12, after=8 if level == 1 else 6)
    style_name = f"Heading {level}"
    p.style = doc.styles[style_name]
    run = p.add_run(text)
    set_run_font(run, size=16 if level == 1 else 13, bold=True, color=HEADING_BLUE if level <= 2 else DARK_BLUE)
    return p


def add_body(doc, text, after=6):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after)
    run = p.add_run(text)
    set_run_font(run, size=11, color=BLACK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, after=8, line=1.167)
    for run in p.runs:
        set_run_font(run, size=11)
    if not p.runs:
        run = p.add_run(text)
    else:
        p.runs[0].text = text
        run = p.runs[0]
    set_run_font(run, size=11, color=BLACK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph_spacing(p, after=8, line=1.167)
    run = p.add_run(text)
    set_run_font(run, size=11, color=BLACK)
    return p


def add_metadata_table(doc):
    rows = [
        ("整理日期", "2026-06-03"),
        ("录音文件", "/Users/linchenhao/Downloads/家长因孩子被打追责对方.m4a"),
        ("录音时长", "约 87 分钟"),
        ("整理说明", "本纪要根据本地语音转写整理。录音中多人交叉说话、情绪激烈，个别人名和个别细节建议结合原录音、监控和书面材料复核。"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_cell_margins(table)
    set_table_width(table, [1680, 7680])
    for i, (label, value) in enumerate(rows):
        label_cell, value_cell = table.rows[i].cells
        shade_cell(label_cell, LIGHT_GRAY)
        for cell_text, cell, bold in [(label, label_cell, True), (value, value_cell, False)]:
            p = cell.paragraphs[0]
            p.text = ""
            run = p.add_run(cell_text)
            set_run_font(run, size=10.5, bold=bold, color=BLACK)
    doc.add_paragraph()


def add_action_table(doc):
    rows = [
        ("1", "被打方家长", "提供医院就诊记录、检查报告等材料，并继续观察孩子身体和心理状态。"),
        ("2", "被打方家长", "根据孩子后续情况决定是否申请伤情鉴定。"),
        ("3", "公安机关", "继续依法开展询问、材料收集、伤情判断及后续处理。"),
        ("4", "园方", "保存并按程序提供监控及现场相关材料。"),
        ("5", "园方", "复盘并完善家长入园、家长接触幼儿、突发冲突上报和家校通知机制。"),
        ("6", "园方", "根据家长意愿安排心理支持，避免未经同意让涉事方再次接触被打孩子。"),
        ("7", "涉事家长", "配合公安及园方后续处理，并对自己的孩子做好正向引导，避免事件继续影响孩子间关系。"),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=3)
    table.style = "Table Grid"
    set_cell_margins(table, top=100, bottom=100, start=120, end=120)
    set_table_width(table, [640, 1700, 7020])
    headers = ["序号", "责任方", "事项"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        shade_cell(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.text = ""
        run = p.add_run(header)
        set_run_font(run, size=10.5, bold=True, color=BLACK)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            p = cell.paragraphs[0]
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx < 2 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, size=10.5, color=BLACK)


def configure_styles(doc):
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, side, Inches(1.0))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    for style_name in ["Normal", "List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.10

    for level, size, color in [(1, 16, HEADING_BLUE), (2, 13, HEADING_BLUE), (3, 12, DARK_BLUE)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(16 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(8 if level == 1 else 6)
        style.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(header, after=0)
    run = header.add_run("会议纪要 | 幼儿园冲突处理会议")
    set_run_font(run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer, after=0)
    run = footer.add_run("本纪要不替代公安笔录、医疗诊断或正式调查报告")
    set_run_font(run, size=8.5, color=MUTED)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=16, after=4)
    run = p.add_run("会议纪要")
    set_run_font(run, size=23, bold=True, color=BLACK)

    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=16)
    run = p.add_run("家长因孩子被打追责对方")
    set_run_font(run, size=14, color=MUTED)

    add_metadata_table(doc)

    add_heading(doc, "一、会议主题", 1)
    add_body(doc, "围绕幼儿园内一名幼儿被另一名幼儿家长掌掴一事，家长、园方、公安/社区及相关代表进行情况说明、责任沟通、后续追责和孩子身心安抚安排。")

    add_heading(doc, "二、参会/发言角色", 1)
    for item in ["被打孩子家长。", "涉事打人方家长。", "幼儿园园方负责人及老师。", "公安/社区调解人员。", "家委/律师代表。"]:
        add_bullet(doc, item)

    add_heading(doc, "三、事件背景与主要事实", 1)
    facts = [
        "录音中多次提及，事发时间约为前一日 8:52 左右。",
        "涉事家长到幼儿园，就孩子之间此前矛盾或说法向被打孩子了解/对质。",
        "过程中，涉事家长在老师反应不及的情况下对被打孩子打了一巴掌；老师随后将孩子拉开并保护在身后。",
        "被打孩子家长认为，两个成年人直接找幼儿园孩子对质，本身已对孩子造成压迫和惊吓；动手掌掴更是严重错误。",
        "园方称事发后先安抚孩子和现场情绪，后续约 10:30 联系被打孩子家长。被打方家长认为通知不及时，并质疑为何让涉事家长先行离开。",
        "事件发生后，被打孩子家长报警并要求追责、查看监控、明确园方责任及对方责任。",
    ]
    for item in facts:
        add_number(doc, item)

    add_heading(doc, "四、各方主要意见", 1)
    sections = [
        ("1. 被打孩子家长", [
            "强烈质疑涉事家长为何不通过老师或园方沟通，而是直接找幼儿园孩子对质。",
            "认为即使没有打巴掌，成年人逼问孩子也会造成心理压迫；打巴掌后果更严重。",
            "对园方通知时效、现场拦阻、允许涉事家长离开等处理方式不满。",
            "表示可以接受对方“道歉是应当的”，但不接受调解或谅解。",
            "明确要求依法、秉公处理，认为自己没有额外索赔或陷害对方，只是要求对方为成年人行为承担相应后果。",
            "表示会继续观察孩子身体和心理状态，必要时进一步就医或走伤情鉴定。",
            "不希望涉事家长再去看孩子或当面反复强化事件，以免造成二次刺激。",
        ]),
        ("2. 涉事打人方家长", [
            "承认事情处理不当，并向被打孩子及家长道歉。",
            "解释本意是了解孩子之间此前矛盾或说法，希望两个孩子把事情说清楚。",
            "表示当时情绪激动，动作发生太快，未能控制住。",
            "希望被打方考虑对自己孩子后续生活、学习或心理影响，给一定余地。",
            "被打方家长对此回应称，不能把涉事家长应承担的后果转化为对受害方的道德压力。",
        ]),
        ("3. 园方", [
            "承认事件发生在幼儿园内，园方管理存在漏洞，不推卸、不逃避。",
            "说明事发后第一时间重点是稳定孩子情绪和现场秩序。",
            "表示已支持被打方报警、调取监控和依法追责。",
            "表示已向教育局、社区等报备，并已向全体家长作出说明。",
            "愿意再次向被打孩子及家长郑重道歉。",
            "提出可根据需要为孩子提供心理支持，优先考虑孩子熟悉的心理相关人员，减少陌生干预带来的不适。",
        ]),
        ("4. 公安/社区调解人员", [
            "认可被打方家长作为父母的情绪可以理解，但建议会议中尽量控制情绪，以便解决问题。",
            "明确指出涉事家长直接找幼儿园孩子了解情况、并动手打孩子的行为是错误的。",
            "说明后续公安处理流程：收集医院材料、了解孩子与现场情况、必要时做公安司法伤情鉴定；若达到轻伤标准，可能进入刑事程序；未达到则依法按治安程序处理。公安可组织调解，但是否接受调解完全尊重双方自愿。",
            "提醒应优先关注孩子身体和心理状况。",
        ]),
        ("5. 家委/律师代表", [
            "认为公安已介入，后续应在查明事实基础上依法、公正处理。",
            "调解或依法处理都是解决问题的方式，但必须尊重当事方意愿。",
            "肯定园方已进行公开说明和复盘，但重点仍应放在孩子身心健康和后续制度完善上。",
        ]),
    ]
    for title, bullets in sections:
        add_heading(doc, title, 2)
        for item in bullets:
            add_bullet(doc, item)

    add_heading(doc, "五、争议焦点", 1)
    for item in [
        "涉事家长道歉是否足以解决问题。",
        "被打方是否接受调解或出具谅解。",
        "是否需要做伤情鉴定，以及伤情对后续处罚性质的影响。",
        "园方是否存在通知不及时、现场处置不充分、管理制度有漏洞等责任。",
        "是否允许涉事家长再次见被打孩子并道歉。",
        "如何避免对孩子造成二次心理刺激。",
    ]:
        add_number(doc, item)

    add_heading(doc, "六、会议形成的初步结论", 1)
    for item in [
        "涉事家长行为不当，需为直接找幼儿园孩子对质及掌掴行为承担责任。",
        "被打方家长接受“道歉是应当的”这一点，但明确不接受调解/谅解，主张依法处理。",
        "公安将按程序收集证据、查看伤情材料，并依法作出处置；调解仅在双方自愿基础上进行。",
        "园方承认管理存在漏洞，承诺配合调查、监控调取、上报和后续整改。",
        "当前优先事项是观察孩子身体伤情和心理状态，避免反复刺激孩子。",
    ]:
        add_number(doc, item)

    add_heading(doc, "七、后续待办事项", 1)
    add_action_table(doc)

    add_heading(doc, "八、备注", 1)
    add_body(doc, "本纪要不替代公安笔录、医疗诊断、伤情鉴定或园方正式调查报告。涉及责任认定、处罚结果、伤情等级等事项，以公安机关、医疗机构和相关主管部门正式文件为准。")

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
